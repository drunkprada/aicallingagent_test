from pathlib import Path

import openpyxl
import pytest

from src.models.schema import Candidate, CandidateStatus
from src.storage import excel_export
from src.storage.db import init_db, upsert_candidate


@pytest.fixture
def stored_candidates(tmp_path, monkeypatch):
    from src.storage.db import get_all_candidates

    db_path = tmp_path / "test.db"
    init_db(db_path)
    monkeypatch.setattr(excel_export, "get_all_candidates", lambda: get_all_candidates(db_path))
    return db_path, tmp_path / "exports" / "candidates.xlsx"


def test_export_creates_expected_columns(tmp_path: Path, monkeypatch):
    db_path = tmp_path / "test.db"
    init_db(db_path)
    monkeypatch.setattr("src.storage.excel_export.get_all_candidates", lambda: [
        Candidate(
            resume_hash="h1",
            source_filename="a.pdf",
            status=CandidateStatus.EXTRACTED,
            full_name="Alice",
            skills=["Go", "Rust"],
        )
    ])

    output_path = tmp_path / "candidates.xlsx"
    result_path = excel_export.export_candidates(output_path)

    assert result_path == output_path
    assert output_path.exists()

    workbook = openpyxl.load_workbook(output_path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert "Name" in headers
    assert "Skills" in headers

    name_col = headers.index("Name") + 1
    skills_col = headers.index("Skills") + 1
    assert sheet.cell(row=2, column=name_col).value == "Alice"
    assert sheet.cell(row=2, column=skills_col).value == "Go, Rust"
    workbook.close()


def test_export_preserves_literal_text_and_numeric_types(stored_candidates):
    db_path, output = stored_candidates
    text_fields = {
        "full_name": "=1+1", "email": "=2+2", "phone": "+441234567890",
        "current_title": "#N/A", "education": "=3+3", "location": "@Edinburgh",
        "current_salary": "-500", "salary_expectation": "=4+4",
        "nationality": "#VALUE!", "notice_period": "=5+5",
        "call_status": "#REF!", "last_call_at": "=6+6",
        "extraction_notes": "=7+7", "source_filename": "=resume.docx",
        "created_at": "=8+8",
    }
    candidate = Candidate(
        resume_hash="literal", skills=["=1+1", "C++"], years_experience=2.5,
        call_attempts=3, status=CandidateStatus.NEEDS_REVIEW, **text_fields,
    )
    upsert_candidate(candidate, db_path)

    assert excel_export.export_candidates(output) == output

    workbook = openpyxl.load_workbook(output, data_only=False)
    try:
        sheet = workbook.active
        values = {**text_fields, "skills": "=1+1, C++", "status": "needs_review",
                  "years_experience": 2.5, "call_attempts": 3}
        assert sheet.max_row == 2
        for col, (field, _) in enumerate(excel_export.COLUMNS, start=1):
            cell = sheet.cell(2, col)
            assert cell.value == values[field]
            assert cell.data_type == ("s" if isinstance(values[field], str) else "n")
        assert sheet["A1"].font.bold
    finally:
        workbook.close()


@pytest.mark.parametrize("invalid", ["\x00", "\x0b", "\x1f", "\ufffe", "\uffff"])
def test_invalid_xml_text_does_not_break_export_or_change_database(stored_candidates, invalid):
    from src.storage.db import get_candidate

    db_path, output = stored_candidates
    text = f"Jane{invalid}Doe\t\nCafé — 東京"
    upsert_candidate(Candidate(
        resume_hash="invalid-text", source_filename="resume.docx",
        full_name=text, skills=[f"Python{invalid}SQL"],
    ), db_path)
    upsert_candidate(Candidate(
        resume_hash="ordinary", source_filename="ordinary.docx", full_name="Alice",
    ), db_path)

    excel_export.export_candidates(output)

    workbook = openpyxl.load_workbook(output)
    try:
        rows = list(workbook.active.values)
        assert len(rows) == 3
        assert rows[1][0] == text.replace(invalid, "\ufffd")
        assert rows[1][5] == "Python\ufffdSQL"
        assert rows[2][0] == "Alice"
        # The export adapts XML-incompatible text, never the source record.
        assert get_candidate("invalid-text", db_path).full_name == text
    finally:
        workbook.close()


@pytest.mark.parametrize("failure", ["save", "replace"])
def test_failed_refresh_preserves_previous_workbook_and_can_retry(
    stored_candidates, monkeypatch, failure
):
    db_path, output = stored_candidates
    candidate = Candidate(resume_hash="h1", source_filename="resume.docx", full_name="Before")
    upsert_candidate(candidate, db_path)
    excel_export.export_candidates(output)
    previous_bytes = output.read_bytes()
    candidate.full_name = "After"
    upsert_candidate(candidate, db_path)

    def broken_save(self, path):
        Path(path).write_bytes(b"incomplete ZIP")
        raise OSError("simulated disk failure")

    def locked_destination(self, target):
        raise PermissionError("simulated locked workbook")

    with monkeypatch.context() as patch:
        if failure == "save":
            patch.setattr(excel_export.Workbook, "save", broken_save)
        else:
            patch.setattr(Path, "replace", locked_destination)
        with pytest.raises(OSError):
            excel_export.export_candidates(output)

    assert output.read_bytes() == previous_bytes
    assert list(output.parent.iterdir()) == [output]
    # Stored results survive an export error; regeneration needs no re-ingestion.
    excel_export.export_candidates(output)
    workbook = openpyxl.load_workbook(output)
    try:
        assert workbook.active["A2"].value == "After"
    finally:
        workbook.close()
    assert list(output.parent.iterdir()) == [output]


def test_failed_first_export_leaves_no_incomplete_file(stored_candidates, monkeypatch):
    _, output = stored_candidates

    def broken_save(self, path):
        Path(path).write_bytes(b"incomplete ZIP")
        raise OSError("simulated disk failure")

    monkeypatch.setattr(excel_export.Workbook, "save", broken_save)
    with pytest.raises(OSError):
        excel_export.export_candidates(output)

    assert not output.exists()
    assert list(output.parent.iterdir()) == []
