import shutil
from unittest.mock import MagicMock

import httpx
import openpyxl
import pytest
from docx import Document
from openai import OpenAI, OpenAIError

from src import main
from src.ingestion import resume_loader
from src.ingestion.extractor import ExtractionError, extract_candidate
from src.models.schema import Candidate, CandidateStatus, ResumeExtraction
from src.storage import db, excel_export


def test_invalid_extraction_mode_fails_before_side_effects(monkeypatch):
    ensure_directories = MagicMock()
    init_db = MagicMock()
    openai_client = MagicMock()

    monkeypatch.setattr(main, "EXTRACTION_MODE", "typo")
    monkeypatch.setattr(main, "ensure_directories", ensure_directories)
    monkeypatch.setattr(main, "init_db", init_db)
    monkeypatch.setattr(main.openai, "OpenAI", openai_client)

    with pytest.raises(
        ValueError,
        match="Invalid EXTRACTION_MODE.*'openai' or 'mock'.*'typo'",
    ):
        main.run_ingest()

    ensure_directories.assert_not_called()
    init_db.assert_not_called()
    openai_client.assert_not_called()


@pytest.fixture
def offline_ingest(tmp_path, monkeypatch):
    incoming = tmp_path / "incoming"
    processed = tmp_path / "processed"
    failed = processed / "failed"
    for directory in (incoming, processed, failed):
        directory.mkdir(parents=True, exist_ok=True)
    db_path = tmp_path / "candidates.db"
    output = tmp_path / "candidates.xlsx"
    extractor = MagicMock()
    client = MagicMock()
    client.return_value.__enter__.return_value = client.return_value

    monkeypatch.setattr(resume_loader, "INCOMING_RESUMES_DIR", incoming)
    monkeypatch.setattr(resume_loader, "PROCESSED_DIR", processed)
    monkeypatch.setattr(resume_loader, "FAILED_DIR", failed)
    monkeypatch.setattr(main, "EXTRACTION_MODE", "openai")
    monkeypatch.setattr(main, "ensure_directories", lambda: None)
    monkeypatch.setattr(main, "init_db", lambda: db.init_db(db_path))
    monkeypatch.setattr(main, "get_candidate", lambda h: db.get_candidate(h, db_path))
    monkeypatch.setattr(main, "upsert_candidate", lambda c: db.upsert_candidate(c, db_path))
    monkeypatch.setattr(excel_export, "get_all_candidates", lambda: db.get_all_candidates(db_path))
    monkeypatch.setattr(main, "export_candidates", lambda: excel_export.export_candidates(output))
    monkeypatch.setattr(main.openai, "OpenAI", client)
    monkeypatch.setattr(main, "extract_candidate", extractor)
    return incoming, processed, failed, db_path, output, extractor, client


def _resume(incoming, name="resume.docx"):
    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("jane@example.com")
    path = incoming / name
    document.save(path)
    return path, resume_loader.hash_file(path)


def test_failed_extraction_can_be_retried(offline_ingest):
    incoming, processed, failed, db_path, _, extractor, _ = offline_ingest
    path, resume_hash = _resume(incoming)
    extractor.side_effect = [
        ExtractionError("temporary API failure"),
        ResumeExtraction(full_name="Jane Doe", email="jane@example.com"),
    ]

    main.run_ingest()
    first = db.get_candidate(resume_hash, db_path)
    assert (failed / path.name).exists()
    shutil.move(failed / path.name, incoming / path.name)

    main.run_ingest()
    recovered = db.get_candidate(resume_hash, db_path)
    assert recovered.full_name == "Jane Doe"
    assert recovered.status == CandidateStatus.EXTRACTED
    assert recovered.extraction_notes is None
    assert recovered.created_at == first.created_at
    assert extractor.call_count == 2
    assert len(db.get_all_candidates(db_path)) == 1
    assert (processed / path.name).exists()
    assert not (failed / path.name).exists()


@pytest.mark.parametrize("status", [
    CandidateStatus.EXTRACTED, CandidateStatus.NEEDS_REVIEW, CandidateStatus.CALL_COMPLETED,
])
def test_successful_records_are_still_deduplicated(offline_ingest, status):
    incoming, processed, _, db_path, _, extractor, client = offline_ingest
    path, resume_hash = _resume(incoming)
    db.init_db(db_path)
    db.upsert_candidate(Candidate(
        resume_hash=resume_hash, source_filename="original.docx", status=status,
        full_name="Reviewed Name", extraction_notes="missing both email and phone",
        call_attempts=2, salary_expectation="reviewed value",
    ), db_path)
    before = db.get_candidate(resume_hash, db_path)

    main.run_ingest()

    extractor.assert_not_called()
    client.assert_not_called()
    assert db.get_candidate(resume_hash, db_path) == before
    assert (processed / path.name).exists()
    assert len(db.get_all_candidates(db_path)) == 1


def test_legacy_failure_retry_preserves_existing_call_data(offline_ingest):
    incoming, processed, _, db_path, _, extractor, _ = offline_ingest
    path, resume_hash = _resume(incoming)
    db.init_db(db_path)
    db.upsert_candidate(Candidate(
        resume_hash=resume_hash, source_filename="old-name.docx",
        status=CandidateStatus.NEEDS_REVIEW,
        extraction_notes="extraction error: previous API failure",
        salary_expectation="45000", nationality="provided answer",
        notice_period="one month", call_status="failed", call_attempts=2,
        last_call_at="2026-01-01T10:00:00+00:00",
    ), db_path)
    before = db.get_candidate(resume_hash, db_path)
    extractor.return_value = ResumeExtraction(full_name="Jane Doe", email="jane@example.com")

    main.run_ingest()

    after = db.get_candidate(resume_hash, db_path)
    assert after.full_name == "Jane Doe"
    assert after.source_filename == path.name
    assert after.status == CandidateStatus.EXTRACTED
    assert after.extraction_notes is None
    for field in ("created_at", "salary_expectation", "nationality", "notice_period",
                  "call_status", "call_attempts", "last_call_at"):
        assert getattr(after, field) == getattr(before, field)
    assert (processed / path.name).exists()


def test_extraction_failure_does_not_stop_batch_or_disappear_from_export(offline_ingest):
    incoming, processed, failed, db_path, output, extractor, _ = offline_ingest
    bad, bad_hash = _resume(incoming, "a-failure.docx")
    good, _ = _resume(incoming, "b-success.docx")
    # Make the two resumes different even if ZIP timestamps are identical.
    document = Document(good)
    document.add_paragraph("Python")
    document.save(good)
    good_hash = resume_loader.hash_file(good)
    extractor.side_effect = [
        ExtractionError("temporary API failure"),
        ResumeExtraction(full_name="Jane Doe", email="jane@example.com"),
    ]

    main.run_ingest()

    assert db.get_candidate(bad_hash, db_path).status.value == "extraction_failed"
    assert db.get_candidate(good_hash, db_path).status == CandidateStatus.EXTRACTED
    assert (failed / bad.name).exists()
    assert (processed / good.name).exists()
    workbook = openpyxl.load_workbook(output)
    rows = list(workbook.active.values)
    status_column = rows[0].index("Status")
    assert {row[status_column] for row in rows[1:]} == {"extraction_failed", "extracted"}
    workbook.close()


def test_repeated_extraction_failure_preserves_review_and_call_data(offline_ingest):
    incoming, _, failed, db_path, _, extractor, _ = offline_ingest
    path, resume_hash = _resume(incoming)
    db.init_db(db_path)
    db.upsert_candidate(Candidate(
        resume_hash=resume_hash, source_filename=path.name,
        status=CandidateStatus.NEEDS_REVIEW,
        extraction_notes="extraction error: old failure", full_name="Reviewed Name",
        call_attempts=2, call_status="failed",
    ), db_path)
    before = db.get_candidate(resume_hash, db_path)
    extractor.side_effect = ExtractionError("new failure")

    main.run_ingest()

    after = db.get_candidate(resume_hash, db_path)
    assert after.status.value == "extraction_failed"
    assert after.extraction_notes == "extraction error: new failure"
    assert after.full_name == before.full_name
    assert after.call_attempts == before.call_attempts
    assert after.call_status == before.call_status
    assert after.created_at == before.created_at
    assert (failed / path.name).exists()


def test_retry_in_mock_mode_stays_flagged_for_review(offline_ingest, monkeypatch):
    incoming, processed, _, db_path, _, extractor, client = offline_ingest
    path, resume_hash = _resume(incoming)
    db.init_db(db_path)
    db.upsert_candidate(Candidate(
        resume_hash=resume_hash, source_filename=path.name,
        status=CandidateStatus.NEEDS_REVIEW,
        extraction_notes="extraction error: old failure",
    ), db_path)
    monkeypatch.setattr(main, "EXTRACTION_MODE", "mock")

    main.run_ingest()

    after = db.get_candidate(resume_hash, db_path)
    assert after.full_name == "Jane Doe"
    assert after.status == CandidateStatus.NEEDS_REVIEW
    assert after.extraction_notes.startswith(main.MOCK_MODE_NOTE)
    assert (processed / path.name).exists()
    client.assert_not_called()
    extractor.assert_not_called()


def test_empty_batch_exports_without_opening_an_api_client(offline_ingest):
    _, _, _, db_path, output, extractor, client = offline_ingest

    main.run_ingest()

    client.assert_not_called()
    extractor.assert_not_called()
    assert db.get_all_candidates(db_path) == []
    workbook = openpyxl.load_workbook(output)
    try:
        assert workbook.active.max_row == 1
    finally:
        workbook.close()


def test_malformed_batch_is_archived_and_exported_without_api_client(offline_ingest):
    incoming, _, failed, db_path, output, extractor, client = offline_ingest
    for filename in ("bad.docx", "bad.pdf"):
        (incoming / filename).write_bytes(b"not a resume document")

    main.run_ingest()

    client.assert_not_called()
    extractor.assert_not_called()
    assert (failed / "bad.docx").exists()
    assert (failed / "bad.pdf").exists()
    assert db.get_all_candidates(db_path) == []
    assert output.exists()


def test_api_client_opens_after_parsing_and_is_reused_for_the_batch(offline_ingest, monkeypatch):
    incoming, _, _, _, _, extractor, client = offline_ingest
    first, _ = _resume(incoming, "a-first.docx")
    second, _ = _resume(incoming, "b-second.docx")
    document = Document(second)
    document.add_paragraph("Python")
    document.save(second)
    parse = main.extract_text
    parsed_paths = []

    def observed_parse(path):
        # The first valid document must be parsed before credentials are needed.
        assert client.call_count == len(parsed_paths)
        parsed_paths.append(path)
        return parse(path)

    monkeypatch.setattr(main, "extract_text", observed_parse)
    extractor.return_value = ResumeExtraction(full_name="Jane Doe", email="jane@example.com")

    main.run_ingest()

    assert parsed_paths == [first, second]
    client.assert_called_once_with()
    client.return_value.__enter__.assert_called_once()
    client.return_value.__exit__.assert_called_once()
    assert extractor.call_count == 2
    assert all(call.args[1] is client.return_value for call in extractor.call_args_list)


def test_missing_credentials_leave_valid_resume_available_for_retry(offline_ingest, monkeypatch):
    incoming, _, failed, db_path, _, extractor, _ = offline_ingest
    path, _ = _resume(incoming)
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setattr(main.openai, "OpenAI", OpenAI)
    parse = MagicMock(wraps=main.extract_text)
    monkeypatch.setattr(main, "extract_text", parse)

    with pytest.raises(OpenAIError, match="api_key"):
        main.run_ingest()

    parse.assert_called_once_with(path)
    extractor.assert_not_called()
    assert path.exists()
    assert not (failed / path.name).exists()
    assert db.get_all_candidates(db_path) == []


@pytest.mark.parametrize("outcome", ["success", "api_error", "storage_error"])
def test_real_sdk_transport_is_closed_on_success_and_failure(offline_ingest, monkeypatch, outcome):
    incoming, _, _, _, _, _, factory = offline_ingest
    _resume(incoming)
    requests = []

    def respond(request):
        requests.append(request)
        if outcome == "api_error":
            return httpx.Response(400, json={"error": {"message": "offline rejection"}})
        return httpx.Response(200, json={
            "id": "offline-completion", "object": "chat.completion", "created": 0,
            "model": "gpt-4o", "choices": [{"index": 0, "finish_reason": "stop",
                "message": {"role": "assistant", "content": '{"full_name":"Jane Doe"}'}}],
        })

    http_client = httpx.Client(transport=httpx.MockTransport(respond))
    sdk_client = OpenAI(api_key="offline-test-key", http_client=http_client)
    factory.return_value = sdk_client
    monkeypatch.setattr(main, "extract_candidate", extract_candidate)
    if outcome == "storage_error":
        monkeypatch.setattr(main, "upsert_candidate", MagicMock(side_effect=OSError("offline disk error")))
    try:
        if outcome == "storage_error":
            with pytest.raises(OSError, match="offline disk error"):
                main.run_ingest()
        else:
            main.run_ingest()
        assert http_client.is_closed
        assert len(requests) == 1
    finally:
        sdk_client.close()


def test_table_resume_reaches_database_and_export_without_live_extraction(
    offline_ingest, monkeypatch
):
    incoming, processed, failed, db_path, output, extractor, client = offline_ingest
    monkeypatch.setattr(main, "EXTRACTION_MODE", "mock")
    document = Document()
    table = document.add_table(rows=4, cols=1)
    for row, text in zip(table.rows, ["Jane Doe", "jane@example.com", "Skills", "Python, SQL"]):
        row.cells[0].text = text
    table_path = incoming / "a-table.docx"
    document.save(table_path)
    table_hash = resume_loader.hash_file(table_path)
    # A malformed file and a normal resume in the same batch must still follow
    # their existing failed/processed paths; no network client may be created.
    (incoming / "b-corrupt.docx").write_bytes(b"not a DOCX package")
    normal_path, _ = _resume(incoming, "c-normal.docx")

    main.run_ingest()

    candidate = db.get_candidate(table_hash, db_path)
    assert candidate is not None
    assert candidate.full_name == "Jane Doe"
    assert candidate.email == "jane@example.com"
    assert candidate.skills == ["Python", "SQL"]
    assert candidate.status == CandidateStatus.NEEDS_REVIEW
    assert len(db.get_all_candidates(db_path)) == 2
    assert (processed / table_path.name).exists()
    assert (processed / normal_path.name).exists()
    assert (failed / "b-corrupt.docx").exists()
    assert not (failed / table_path.name).exists()

    workbook = openpyxl.load_workbook(output)
    try:
        rows = list(workbook.active.values)
        headers = rows[0]
        table_row = next(row for row in rows[1:] if row[headers.index("Source File")] == table_path.name)
        assert table_row[headers.index("Email")] == "jane@example.com"
        assert table_row[headers.index("Skills")] == "Python, SQL"
        assert table_row[headers.index("Status")] == "needs_review"
    finally:
        workbook.close()
    client.assert_not_called()
    extractor.assert_not_called()


def test_mock_sections_reach_storage_and_export_without_work_history(
    offline_ingest, monkeypatch
):
    incoming, processed, _, db_path, output, extractor, client = offline_ingest
    monkeypatch.setattr(main, "EXTRACTION_MODE", "mock")
    document = Document()
    for line in [
        "Jane Doe", "jane@example.com", "Skills developed through work",
        "This sentence is not a heading", "Technical Skills:",
        "- scikit-learn, Python", "SQL", "Docker", "Git", "Linux", "React",
        "Work Experience", "Built a recruitment portal", "Education",
        "BSc Computer Science", "Example University", "Projects", "A web portal",
    ]:
        document.add_paragraph(line)
    path = incoming / "sections.docx"
    document.save(path)
    resume_hash = resume_loader.hash_file(path)

    main.run_ingest()

    candidate = db.get_candidate(resume_hash, db_path)
    expected_skills = ["scikit-learn", "Python", "SQL", "Docker", "Git", "Linux", "React"]
    assert candidate.skills == expected_skills
    assert candidate.education == "BSc Computer Science\nExample University"
    assert candidate.status == CandidateStatus.NEEDS_REVIEW
    assert candidate.extraction_notes.startswith(main.MOCK_MODE_NOTE)
    assert (processed / path.name).exists()
    workbook = openpyxl.load_workbook(output)
    try:
        headers, row = list(workbook.active.values)
        assert row[headers.index("Skills")] == ", ".join(expected_skills)
        assert row[headers.index("Education")] == candidate.education
        assert row[headers.index("Status")] == "needs_review"
    finally:
        workbook.close()
    client.assert_not_called()
    extractor.assert_not_called()


@pytest.mark.parametrize("failure", ["permission", "disappeared"])
def test_unreadable_resume_does_not_stop_batch_or_export(
    offline_ingest, monkeypatch, caplog, failure
):
    incoming, processed, failed, db_path, output, extractor, client = offline_ingest
    bad, _ = _resume(incoming, "a-unreadable.docx")
    good, good_hash = _resume(incoming, "b-readable.docx")
    real_hash = main.hash_file

    def hash_with_file_error(path):
        if path == bad:
            if failure == "disappeared":
                bad.unlink()
                raise FileNotFoundError("resume removed after scan")
            raise PermissionError("resume is locked")
        return real_hash(path)

    monkeypatch.setattr(main, "hash_file", hash_with_file_error)
    extractor.return_value = ResumeExtraction(
        full_name="Jane Doe", email="jane@example.com"
    )

    main.run_ingest()

    assert bad.exists() is (failure == "permission")
    assert not (processed / bad.name).exists()
    assert not (failed / bad.name).exists()
    assert (processed / good.name).exists()
    candidate = db.get_candidate(good_hash, db_path)
    assert candidate.source_filename == good.name
    assert candidate.status == CandidateStatus.EXTRACTED
    assert len(db.get_all_candidates(db_path)) == 1
    extractor.assert_called_once()
    client.return_value.__exit__.assert_called_once()
    assert "Skipping unreadable resume a-unreadable.docx" in caplog.text
    workbook = openpyxl.load_workbook(output)
    try:
        rows = list(workbook.active.values)
        source_column = rows[0].index("Source File")
        assert [row[source_column] for row in rows[1:]] == [good.name]
    finally:
        workbook.close()


def test_locked_resume_stays_available_and_succeeds_after_unlock(
    offline_ingest, monkeypatch
):
    incoming, processed, failed, db_path, output, extractor, client = offline_ingest
    path, resume_hash = _resume(incoming)
    real_hash = main.hash_file
    monkeypatch.setattr(
        main, "hash_file", MagicMock(side_effect=PermissionError("resume is locked"))
    )

    main.run_ingest()

    assert path.exists()
    assert db.get_all_candidates(db_path) == []
    extractor.assert_not_called()
    client.assert_not_called()
    workbook = openpyxl.load_workbook(output)
    try:
        assert workbook.active.max_row == 1
    finally:
        workbook.close()

    monkeypatch.setattr(main, "hash_file", real_hash)
    extractor.return_value = ResumeExtraction(
        full_name="Jane Doe", email="jane@example.com"
    )
    main.run_ingest()

    assert not path.exists()
    assert (processed / path.name).exists()
    assert not (failed / path.name).exists()
    assert db.get_candidate(resume_hash, db_path).status == CandidateStatus.EXTRACTED
    extractor.assert_called_once()


@pytest.mark.parametrize("outcome", ["success", "duplicate", "parse_failure", "extraction_failure"])
def test_archive_failure_does_not_stop_remaining_resumes_or_export(
    offline_ingest, monkeypatch, caplog, outcome
):
    incoming, processed, failed, db_path, output, extractor, client = offline_ingest
    blocked, blocked_hash = _resume(incoming, "a-blocked.docx")
    good, _ = _resume(incoming, "b-good.docx")
    document = Document(good)
    document.add_paragraph("A distinct second resume")
    document.save(good)
    good_hash = resume_loader.hash_file(good)
    db.init_db(db_path)
    before = None
    if outcome == "duplicate":
        db.upsert_candidate(Candidate(
            resume_hash=blocked_hash, source_filename="original.docx",
            full_name="Reviewed Name", call_attempts=2,
            salary_expectation="Reviewed answer",
        ), db_path)
        before = db.get_candidate(blocked_hash, db_path)
    elif outcome == "parse_failure":
        blocked.write_bytes(b"malformed DOCX")
        blocked_hash = resume_loader.hash_file(blocked)
    extraction = ResumeExtraction(full_name="Jane Doe", email="jane@example.com")
    extractor.side_effect = (
        [ExtractionError("offline rejection"), extraction]
        if outcome == "extraction_failure" else None
    )
    extractor.return_value = extraction

    archive_name = (
        "move_to_failed" if outcome in {"parse_failure", "extraction_failure"}
        else "move_to_processed"
    )
    real_archive = getattr(main, archive_name)

    def locked_archive(path):
        if path == blocked:
            raise PermissionError("archive destination locked")
        return real_archive(path)

    monkeypatch.setattr(main, archive_name, locked_archive)

    main.run_ingest()

    assert blocked.exists()
    assert not (processed / blocked.name).exists()
    assert not (failed / blocked.name).exists()
    assert (processed / good.name).exists()
    assert db.get_candidate(good_hash, db_path).status == CandidateStatus.EXTRACTED
    record = db.get_candidate(blocked_hash, db_path)
    if outcome == "parse_failure":
        assert record is None
    elif outcome == "duplicate":
        assert record == before
    else:
        assert record.status == (
            CandidateStatus.EXTRACTION_FAILED
            if outcome == "extraction_failure" else CandidateStatus.EXTRACTED
        )
    assert "Could not archive a-blocked.docx" in caplog.text
    assert "archive destination locked" in caplog.text
    client.return_value.__exit__.assert_called_once()
    workbook = openpyxl.load_workbook(output)
    try:
        rows = list(workbook.active.values)
        assert len(rows) == len(db.get_all_candidates(db_path)) + 1
        assert "b-good.docx" in {row[rows[0].index("Source File")] for row in rows[1:]}
    finally:
        workbook.close()


def test_successful_extraction_is_not_repeated_after_archive_failure(
    offline_ingest, monkeypatch
):
    incoming, processed, _, db_path, _, extractor, client = offline_ingest
    path, resume_hash = _resume(incoming)
    real_archive = main.move_to_processed
    monkeypatch.setattr(
        main, "move_to_processed", MagicMock(side_effect=PermissionError("locked"))
    )
    extractor.return_value = ResumeExtraction(full_name="Jane Doe", email="jane@example.com")

    main.run_ingest()
    before = db.get_candidate(resume_hash, db_path)
    assert path.exists()
    extractor.assert_called_once()

    monkeypatch.setattr(main, "move_to_processed", real_archive)
    extractor.reset_mock()
    client.reset_mock()
    main.run_ingest()

    assert (processed / path.name).exists()
    assert not path.exists()
    assert db.get_candidate(resume_hash, db_path) == before
    assert len(db.get_all_candidates(db_path)) == 1
    extractor.assert_not_called()
    client.assert_not_called()


@pytest.mark.parametrize("case", ["blank_fields", "negative_experience", "infinite_experience"])
def test_invalid_sdk_fields_are_reviewed_or_failed_without_stopping_batch(
    offline_ingest, monkeypatch, case
):
    incoming, processed, failed, db_path, output, _, factory = offline_ingest
    first, first_hash = _resume(incoming, "a-invalid-fields.docx")
    second, _ = _resume(incoming, "b-valid-fields.docx")
    document = Document(second)
    document.add_paragraph("A different resume")
    document.save(second)
    second_hash = resume_loader.hash_file(second)
    valid = '{"full_name":"Jane Doe","email":"jane@example.com","years_experience":2.5}'
    if case == "blank_fields":
        invalid = '{"full_name":"  ","email":"  ","phone":"  ","skills":[" ","Python"]}'
    else:
        years = "-1" if case == "negative_experience" else "1e309"
        invalid = '{"full_name":"Jane Doe","email":"jane@example.com","years_experience":' + years + '}'
    responses = iter([invalid, valid])
    requests = []

    def respond(request):
        requests.append(request)
        return httpx.Response(200, json={
            "id": "offline-completion", "object": "chat.completion", "created": 0,
            "model": "gpt-4o", "choices": [{"index": 0, "finish_reason": "stop",
                "message": {"role": "assistant", "content": next(responses)}}],
        })

    http_client = httpx.Client(transport=httpx.MockTransport(respond))
    sdk_client = OpenAI(api_key="offline-test-key", http_client=http_client)
    factory.return_value = sdk_client
    monkeypatch.setattr(main, "extract_candidate", extract_candidate)
    try:
        main.run_ingest()
        first_record = db.get_candidate(first_hash, db_path)
        if case == "blank_fields":
            assert first_record.status == CandidateStatus.NEEDS_REVIEW
            assert first_record.full_name is None
            assert first_record.email is None
            assert first_record.phone is None
            assert first_record.skills == ["Python"]
            assert "missing full_name" in first_record.extraction_notes
            assert "missing both email and phone" in first_record.extraction_notes
            assert (processed / first.name).exists()
        else:
            assert first_record.status == CandidateStatus.EXTRACTION_FAILED
            assert first_record.extraction_notes.startswith("extraction error:")
            assert first_record.years_experience is None
            assert (failed / first.name).exists()
        assert (processed / second.name).exists()
        assert db.get_candidate(second_hash, db_path).years_experience == 2.5
        assert len(requests) == 2
        assert http_client.is_closed
        workbook = openpyxl.load_workbook(output)
        try:
            rows = list(workbook.active.values)
            status_column = rows[0].index("Status")
            assert len(rows) == 3
            assert {row[status_column] for row in rows[1:]} == {
                first_record.status.value, "extracted",
            }
        finally:
            workbook.close()
    finally:
        sdk_client.close()
