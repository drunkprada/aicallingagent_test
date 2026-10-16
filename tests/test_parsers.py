from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.ingestion.parsers import ParserError, extract_text, extract_text_from_docx, extract_text_from_pdf


def test_extract_text_from_docx_real_file(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Software Engineer, 5 years experience")
    path = tmp_path / "resume.docx"
    doc.save(path)

    text = extract_text_from_docx(path)
    assert "Jane Doe" in text
    assert "Software Engineer" in text
    assert text == "Jane Doe\nSoftware Engineer, 5 years experience"


def test_extract_text_from_empty_docx_raises(tmp_path: Path):
    doc = Document()
    path = tmp_path / "empty.docx"
    doc.save(path)

    with pytest.raises(ParserError):
        extract_text_from_docx(path)


def test_table_only_docx_is_read_instead_of_rejected(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=4, cols=1)
    lines = ["Jane Doe", "jane@example.com", "Skills", "Python, SQL"]
    for row, text in zip(table.rows, lines):
        row.cells[0].text = text
    path = tmp_path / "table-resume.DOCX"
    doc.save(path)

    assert extract_text(path) == "\n".join(lines)


def test_docx_paragraphs_and_tables_keep_document_order(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Jane Doe")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Email"
    table.cell(0, 1).text = "jane@example.com"
    doc.add_paragraph("Experience")
    table = doc.add_table(rows=2, cols=2)
    for row, texts in zip(table.rows, [("Engineer", "2024–2026"), ("Python", "SQL")]):
        for cell, text in zip(row.cells, texts):
            cell.text = text
    doc.add_paragraph("Education")
    path = tmp_path / "mixed.docx"
    doc.save(path)

    assert extract_text_from_docx(path).splitlines() == [
        "Jane Doe", "Email", "jane@example.com", "Experience",
        "Engineer", "2024–2026", "Python", "SQL", "Education",
    ]


def test_nested_tables_keep_content_between_cell_paragraphs(tmp_path: Path):
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Before nested table"
    nested = cell.add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Skills"
    nested.cell(0, 1).text = "Python, SQL"
    cell.add_paragraph("After nested table")
    doc.add_paragraph("End of resume")
    path = tmp_path / "nested.docx"
    doc.save(path)

    text = extract_text_from_docx(path)
    assert [line for line in text.splitlines() if line] == [
        "Before nested table", "Skills", "Python, SQL",
        "After nested table", "End of resume",
    ]


@pytest.mark.parametrize("direction", ["horizontal", "vertical"])
def test_merged_cells_are_read_once_without_dropping_repeated_text(tmp_path: Path, direction):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    if direction == "horizontal":
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged heading"
        table.cell(1, 0).text = "Python"
        table.cell(1, 1).text = "Python"
    else:
        table.cell(0, 0).merge(table.cell(1, 0)).text = "Merged heading"
        table.cell(0, 1).text = "Python"
        table.cell(1, 1).text = "Python"
    path = tmp_path / "merged.docx"
    doc.save(path)

    assert extract_text_from_docx(path).splitlines() == ["Merged heading", "Python", "Python"]


def test_blank_table_is_still_rejected(tmp_path: Path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = " \t "
    path = tmp_path / "blank-table.docx"
    doc.save(path)

    with pytest.raises(ParserError, match="No extractable text"):
        extract_text_from_docx(path)


def test_corrupt_docx_raises_parser_error(tmp_path: Path):
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"not a DOCX package")

    with pytest.raises(ParserError, match="Failed to read DOCX") as error:
        extract_text_from_docx(path)
    assert error.value.__cause__ is not None


def test_invalid_table_structure_is_reported_as_parser_error(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Jane Doe")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Skills"
    # A ZIP/XML-readable document can still contain invalid table properties.
    span = table.cell(0, 0)._tc.get_or_add_tcPr().get_or_add_gridSpan()
    span.set(qn("w:val"), "not-a-number")
    path = tmp_path / "invalid-table.docx"
    doc.save(path)

    with pytest.raises(ParserError, match="Failed to read DOCX") as error:
        extract_text_from_docx(path)
    assert isinstance(error.value.__cause__, ValueError)


def test_extract_text_from_pdf_uses_pdfplumber(tmp_path: Path, monkeypatch):
    class FakePage:
        def extract_text(self):
            return "John Smith\nData Analyst"

    class FakePdf:
        pages = [FakePage()]

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

    def fake_open(path):
        return FakePdf()

    monkeypatch.setattr("src.ingestion.parsers.pdfplumber.open", fake_open)

    path = tmp_path / "resume.pdf"
    path.write_bytes(b"%PDF-1.4 fake")
    text = extract_text_from_pdf(path)
    assert "John Smith" in text


def test_extract_text_unsupported_extension(tmp_path: Path):
    path = tmp_path / "resume.txt"
    path.write_text("hello")
    with pytest.raises(ParserError):
        extract_text(path)
