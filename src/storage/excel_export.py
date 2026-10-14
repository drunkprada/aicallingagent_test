import re
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.config import EXCEL_EXPORT_PATH
from src.models.schema import Candidate
from src.storage.db import get_all_candidates

COLUMNS = [
    ("full_name", "Name"),
    ("email", "Email"),
    ("phone", "Phone"),
    ("current_title", "Current Title"),
    ("years_experience", "Years Experience"),
    ("skills", "Skills"),
    ("education", "Education"),
    ("location", "Location"),
    ("current_salary", "Current Salary"),
    ("salary_expectation", "Salary Expectation"),
    ("nationality", "Nationality"),
    ("notice_period", "Notice Period"),
    ("status", "Status"),
    ("call_status", "Call Status"),
    ("call_attempts", "Call Attempts"),
    ("last_call_at", "Last Call At"),
    ("extraction_notes", "Extraction Notes"),
    ("source_filename", "Source File"),
    ("created_at", "Added On"),
]

# XML 1.0 cannot represent these code points in an XLSX worksheet.
# Keep tabs, line breaks and valid Unicode; only adapt the exported copy.
INVALID_XML_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _cell_value(candidate: Candidate, field: str):
    value = getattr(candidate, field)
    if field == "skills":
        return ", ".join(value) if value else ""
    if field == "status":
        return value.value if value else ""
    return value if value is not None else ""


def export_candidates(output_path: Path = EXCEL_EXPORT_PATH) -> Path:
    candidates = get_all_candidates()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Candidates"

    for col_idx, (_, header) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)

    for row_idx, candidate in enumerate(candidates, start=2):
        for col_idx, (field, _) in enumerate(COLUMNS, start=1):
            value = _cell_value(candidate, field)
            if isinstance(value, str):
                value = INVALID_XML_TEXT.sub("\ufffd", value)
            cell = sheet.cell(row=row_idx, column=col_idx, value=value)
            if isinstance(value, str):
                # Resume/answer text is never a formula or an Excel error value.
                cell.data_type = "s"

    for col_idx, (field, header) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = max(
            len(header) + 2, 14
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    # Stage beside the destination so replacement stays on the same filesystem.
    # Close the temporary handle first so saving also works on Windows.
    with NamedTemporaryFile(
        dir=output_path.parent, prefix=f".{output_path.name}.", suffix=".tmp", delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        temporary_path.replace(output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
        workbook.close()
    return output_path
